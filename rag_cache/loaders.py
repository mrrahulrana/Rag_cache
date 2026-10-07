from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import Path

from langchain_text_splitters import RecursiveCharacterTextSplitter

from rag_cache.config import Settings
from rag_cache.textutil import content_hash, normalize_text

SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".png",
    ".jpg",
    ".jpeg",
    ".webp",
    ".tif",
    ".tiff",
    ".bmp",
    ".gif",
    ".docx",
    ".xlsx",
    ".xlsm",
    ".xls",
    ".csv",
    ".pptx",
    ".txt",
    ".md",
    ".html",
    ".htm",
}

_HEADING = re.compile(
    r"^(?:chapter\s+\d+\b|section\s+\d+\b|\d+(?:\.\d+){0,3}\s+\S).{0,80}$",
    re.IGNORECASE,
)


class IngestError(ValueError):
    pass


@dataclass
class TextChunk:
    text: str
    section: str | None
    page_number: int | None
    metadata: dict = field(default_factory=dict)
    content_hash: str = ""

    def __post_init__(self) -> None:
        self.text = self.text.replace("\x00", "").strip()
        self.content_hash = content_hash(self.text)


def parse_file(path: Path, settings: Settings) -> list[TextChunk]:
    extension = path.suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise IngestError(f"Unsupported file type {extension or '(none)'}. Supported types: {supported}")
    parser = {
        ".pdf": _parse_pdf,
        ".docx": _parse_docx,
        ".xlsx": _parse_xlsx,
        ".xlsm": _parse_xlsx,
        ".xls": _parse_xls,
        ".csv": _parse_csv,
        ".pptx": _parse_pptx,
        ".txt": _parse_text,
        ".md": _parse_text,
        ".html": _parse_html,
        ".htm": _parse_html,
    }.get(extension, _parse_image)
    chunks = parser(path, settings)
    return [chunk for chunk in chunks if normalize_text(chunk.text)]


def _splitter(settings: Settings) -> RecursiveCharacterTextSplitter:
    return RecursiveCharacterTextSplitter(
        chunk_size=settings.chunk_size,
        chunk_overlap=settings.chunk_overlap,
        separators=["\n\n", "\n", ". ", " ", ""],
    )


def _split_prose(
    text: str,
    *,
    settings: Settings,
    section: str | None,
    page_number: int | None,
    metadata: dict,
) -> list[TextChunk]:
    splitter = _splitter(settings)
    chunks: list[TextChunk] = []
    for block_section, body in _section_blocks(text, section):
        parts = splitter.split_text(body) or [body]
        for part in parts:
            cleaned = part.replace("\x00", "").strip()
            if len(normalize_text(cleaned)) < 15:
                continue
            chunks.append(
                TextChunk(
                    text=cleaned,
                    section=block_section,
                    page_number=page_number,
                    metadata=dict(metadata),
                )
            )
    return chunks


def _section_blocks(text: str, default_section: str | None) -> list[tuple[str | None, str]]:
    blocks: list[tuple[str | None, str]] = []
    current = default_section
    buffer: list[str] = []

    def flush() -> None:
        body = "\n".join(buffer).strip()
        if body:
            blocks.append((current, body))

    for line in text.splitlines():
        stripped = line.strip()
        if _looks_like_heading(stripped):
            flush()
            buffer = []
            current = stripped
            continue
        buffer.append(line)
    flush()
    if not blocks and text.strip():
        blocks.append((default_section, text.strip()))
    return blocks


def _looks_like_heading(line: str) -> bool:
    if not line or len(line) < 3 or len(line) > 90:
        return False
    if line.endswith((".", ";", ",")):
        return False
    if _HEADING.match(line):
        return True
    words = line.split()
    letters = [char for char in line if char.isalpha()]
    return bool(letters) and len(words) <= 8 and line.upper() == line and any(char.islower() or char.isupper() for char in letters)


def _cell(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _table_chunks(
    rows: list[list[str]],
    *,
    settings: Settings,
    section: str | None,
    page_number: int | None,
    metadata: dict,
) -> list[TextChunk]:
    if not rows or not any(any(cell for cell in row) for row in rows):
        return []
    header = rows[0]
    body = rows[1:]
    windows = [body[index : index + 20] for index in range(0, len(body), 20)] or [[]]
    chunks: list[TextChunk] = []
    for offset, window in enumerate(windows):
        row_start = 2 + offset * 20 if body else 1
        row_end = row_start + len(window) - 1 if window else 1
        lines = [" | ".join(header)]
        lines.extend(" | ".join(row) for row in window)
        text = "\n".join(lines).strip()
        chunk_meta = dict(metadata)
        chunk_meta["row_start"] = row_start
        chunk_meta["row_end"] = row_end
        pieces = _split_prose(
            text,
            settings=settings,
            section=section,
            page_number=page_number,
            metadata=chunk_meta,
        )
        chunks.extend(pieces or [TextChunk(text=text, section=section, page_number=page_number, metadata=chunk_meta)])
    return [chunk for chunk in chunks if normalize_text(chunk.text)]


def _parse_pdf(path: Path, settings: Settings) -> list[TextChunk]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise IngestError("pypdf is not installed.") from exc
    try:
        reader = PdfReader(str(path))
    except Exception as exc:
        raise IngestError(f"Could not read PDF: {exc}") from exc
    if getattr(reader, "is_encrypted", False):
        try:
            reader.decrypt("")
        except Exception as exc:
            raise IngestError("Password-protected PDFs are not supported.") from exc
    chunks: list[TextChunk] = []
    for page_number, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        if not text.strip():
            text = _ocr_pdf_page(page)
        chunks.extend(
            _split_prose(
                text,
                settings=settings,
                section=None,
                page_number=page_number,
                metadata={"source_type": "pdf"},
            )
        )
    return chunks


def _ocr_pdf_page(page: object) -> str:
    images = getattr(page, "images", None)
    if not images:
        return ""
    parts: list[str] = []
    for image in images:
        data = getattr(image, "data", None)
        if data:
            parts.append(_ocr_bytes(data))
    return "\n".join(part for part in parts if part.strip())


def _parse_docx(path: Path, settings: Settings) -> list[TextChunk]:
    try:
        from docx import Document
    except ImportError as exc:
        raise IngestError("python-docx is not installed.") from exc
    try:
        document = Document(str(path))
    except Exception as exc:
        raise IngestError(f"Could not read Word document: {exc}") from exc
    section = None
    blocks: list[tuple[str | None, str]] = []
    buffer: list[str] = []

    def flush() -> None:
        body = "\n".join(buffer).strip()
        if body:
            blocks.append((section, body))
        buffer.clear()

    for paragraph in document.paragraphs:
        style = paragraph.style.name if paragraph.style is not None else ""
        text = paragraph.text.strip()
        if not text:
            continue
        if style.startswith("Heading"):
            flush()
            section = text
            continue
        buffer.append(text)
    flush()
    for table in document.tables:
        rows = [[_cell(cell.text) for cell in row.cells] for row in table.rows]
        blocks.append((section, "\n".join(" | ".join(row) for row in rows)))
    chunks: list[TextChunk] = []
    for block_section, body in blocks:
        chunks.extend(
            _split_prose(
                body,
                settings=settings,
                section=block_section,
                page_number=None,
                metadata={"source_type": "docx"},
            )
        )
    return chunks


def _parse_xlsx(path: Path, settings: Settings) -> list[TextChunk]:
    try:
        import openpyxl
    except ImportError as exc:
        raise IngestError("openpyxl is not installed.") from exc
    try:
        workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    except Exception as exc:
        raise IngestError(f"Could not read Excel workbook: {exc}") from exc
    chunks: list[TextChunk] = []
    try:
        for page_number, sheet in enumerate(workbook.worksheets, start=1):
            rows = [[_cell(value) for value in row] for row in sheet.iter_rows(values_only=True)]
            chunks.extend(
                _table_chunks(
                    rows,
                    settings=settings,
                    section=sheet.title,
                    page_number=page_number,
                    metadata={"source_type": "excel", "sheet_name": sheet.title},
                )
            )
    finally:
        workbook.close()
    return chunks


def _parse_xls(path: Path, settings: Settings) -> list[TextChunk]:
    try:
        import xlrd
    except ImportError as exc:
        raise IngestError("xlrd is not installed.") from exc
    try:
        workbook = xlrd.open_workbook(str(path))
    except Exception as exc:
        raise IngestError(f"Could not read Excel workbook: {exc}") from exc
    chunks: list[TextChunk] = []
    for page_number, sheet in enumerate(workbook.sheets(), start=1):
        rows = []
        for row_index in range(sheet.nrows):
            rows.append([_cell(value) for value in sheet.row_values(row_index)])
        chunks.extend(
            _table_chunks(
                rows,
                settings=settings,
                section=sheet.name,
                page_number=page_number,
                metadata={"source_type": "excel", "sheet_name": sheet.name},
            )
        )
    return chunks


def _parse_csv(path: Path, settings: Settings) -> list[TextChunk]:
    raw = path.read_bytes()
    text = None
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise IngestError("Could not decode this CSV file.")
    rows = [[_cell(value) for value in row] for row in csv.reader(io.StringIO(text))]
    return _table_chunks(
        rows,
        settings=settings,
        section=path.stem,
        page_number=1,
        metadata={"source_type": "csv", "sheet_name": path.stem},
    )


def _parse_pptx(path: Path, settings: Settings) -> list[TextChunk]:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise IngestError("python-pptx is not installed.") from exc
    try:
        presentation = Presentation(str(path))
    except Exception as exc:
        raise IngestError(f"Could not read PowerPoint file: {exc}") from exc
    chunks: list[TextChunk] = []
    for page_number, slide in enumerate(presentation.slides, start=1):
        title = None
        if slide.shapes.title is not None:
            title = slide.shapes.title.text.strip() or None
        lines: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = shape.text_frame.text.strip()
                if text and text != title:
                    lines.append(text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    lines.append(" | ".join(_cell(cell.text) for cell in row.cells))
        chunks.extend(
            _split_prose(
                "\n".join(lines),
                settings=settings,
                section=title,
                page_number=page_number,
                metadata={"source_type": "pptx", "slide": page_number},
            )
        )
    return chunks


def _parse_text(path: Path, settings: Settings) -> list[TextChunk]:
    raw = path.read_bytes()
    text = None
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise IngestError("Could not decode this text file.")
    return _split_prose(
        text,
        settings=settings,
        section=None,
        page_number=1,
        metadata={"source_type": "text"},
    )


class _HTMLText(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self._skip = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in {"script", "style", "noscript"}:
            self._skip += 1
        if tag in {"h1", "h2", "h3", "h4", "p", "div", "li", "tr", "br"}:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in {"script", "style", "noscript"} and self._skip:
            self._skip -= 1
        if tag in {"h1", "h2", "h3", "h4", "p", "li"}:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self._skip:
            self.parts.append(data)


def _parse_html(path: Path, settings: Settings) -> list[TextChunk]:
    parser = _HTMLText()
    parser.feed(path.read_text(encoding="utf-8", errors="ignore"))
    return _split_prose(
        "".join(parser.parts),
        settings=settings,
        section=None,
        page_number=1,
        metadata={"source_type": "html"},
    )


def _parse_image(path: Path, settings: Settings) -> list[TextChunk]:
    text = _ocr_bytes(path.read_bytes())
    if not text.strip():
        raise IngestError(
            "No text could be read from this image. Install Tesseract OCR and make sure it is on PATH."
        )
    return _split_prose(
        text,
        settings=settings,
        section="image",
        page_number=1,
        metadata={"source_type": "image"},
    )


def _ocr_bytes(data: bytes) -> str:
    try:
        import pytesseract
        from PIL import Image
    except ImportError:
        return ""
    _configure_tesseract(pytesseract)
    try:
        with Image.open(io.BytesIO(data)) as image:
            return pytesseract.image_to_string(image) or ""
    except Exception:
        return ""


def _configure_tesseract(pytesseract: object) -> None:
    candidates = [
        Path(r"C:\Program Files\Tesseract-OCR\tesseract.exe"),
        Path(r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe"),
    ]
    for candidate in candidates:
        if candidate.exists():
            pytesseract.pytesseract.tesseract_cmd = str(candidate)
            return
